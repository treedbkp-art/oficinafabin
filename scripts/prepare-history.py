#!/usr/bin/env python3
"""Read original workshop workbooks, without modifying them. Output private import manifests.
Usage: python prepare-history.py INPUT_DIRECTORY OUTPUT_DIRECTORY [--cache RAW_JSON]
Requires openpyxl. Do not commit source workbooks, caches, or generated manifests.
"""
import argparse,calendar,collections,datetime,hashlib,json,re,unicodedata,zipfile
from pathlib import Path
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
import xml.etree.ElementTree as ET
MONTHS=['Janeiro','Fevereiro','Março','Abril','Maio','Junho','Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']
def text(v):
 if v is None:return ''
 if isinstance(v,dict):return v.get('date',json.dumps(v,ensure_ascii=False))
 if isinstance(v,float) and v.is_integer():return str(int(v))
 return str(v).strip()
def norm(v):return ''.join(c for c in unicodedata.normalize('NFD',text(v)) if not unicodedata.combining(c)).lower().strip()
def money(v):
 if v is None or isinstance(v,(dict,bool)):return None
 s=text(v).replace('R$','').replace(' ','')
 if ',' in s:s=s.replace('.','').replace(',','.')
 if not re.fullmatch(r'-?\d+(?:\.\d+)?',s):return None
 try:return int((Decimal(s)*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
 except InvalidOperation:return None
def value(c,a):return c.get(a,{}).get('v')
def read_workbooks(folder):
 from openpyxl import load_workbook
 files=[]
 def encode(v):return {'date':v.isoformat()} if isinstance(v,(datetime.date,datetime.datetime)) else v
 for f in sorted(folder.glob('*.xlsx')):
  if not re.search(r'20\d{2}',f.name):continue
  w=load_workbook(f,read_only=True,data_only=False);cached=load_workbook(f,read_only=True,data_only=True);sheets=[]
  for sheet in w:
   cells={}
   for row,vr in zip(sheet.iter_rows(),cached[sheet.title].iter_rows()):
    for c,vc in zip(row,vr):
     if c.value is None:continue
     e={'v':encode(vc.value)}
     if c.data_type=='f':e['f']=c.value if isinstance(c.value,str) else vars(c.value)
     if c.is_date:e['format']=c.number_format
     cells[c.coordinate]=e
   sheets.append({'name':sheet.title,'cells':cells})
  with zipfile.ZipFile(f) as z:
   ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
   rels={e.attrib['Id']:e.attrib['Target'] for e in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
   for elem,sheet in zip(ET.fromstring(z.read('xl/workbook.xml')).find('m:sheets',ns),sheets):
    target=rels[elem.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
    target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
    for cell in ET.fromstring(z.read(target)).findall('.//m:sheetData/m:row/m:c',ns):
     v=cell.find('m:v',ns);entry=sheet['cells'].get(cell.attrib['r'])
     if entry is not None and v is not None:entry['stored']=v.text
  files.append({'filename':f.name,'year':int(re.search(r'20\d{2}',f.name)[0]),'sheets':sheets});w.close();cached.close()
 return files

def parse_date(v,year,month):
 original=text(v);parsed=None;suggestion=None
 try:
  if isinstance(v,dict) and 'date' in v:parsed=datetime.date.fromisoformat(v['date'][:10])
  elif re.fullmatch(r'\d{1,2}/\d{1,2}/(?:\d{2}|\d{4})',original):
   day,mo,yr=map(int,original.split('/'));parsed=datetime.date(yr+2000 if yr<100 else yr,mo,day)
  elif re.fullmatch(r'\d{4}-\d{2}-\d{2}',original):parsed=datetime.date.fromisoformat(original)
 except ValueError:pass
 if parsed and parsed.year==year and parsed.month==month:return parsed.isoformat(),parsed.isoformat(),'day',None
 if parsed and parsed.year==year and parsed.day==month:
  try:suggestion=datetime.date(year,month,parsed.month).isoformat()
  except ValueError:pass
 # The sheet period is a filing reference, never an invented service day.
 return f'{year}-{month:02}-01',f'{year}-{month:02}-{calendar.monthrange(year,month)[1]:02}','month',suggestion

def prepare(f,folder):
 records=[];sheets=[];excluded=[];unmapped=[]
 for s in f['sheets']:
  c=s['cells'];name=s['name']
  if name not in MONTHS:
   sheets.append({'name':name,'kind':'other','records':0,'nonemptyCells':len(c),'note':'Aba original preservada; não contém fichas de orçamento.'});continue
  month=MONTHS.index(name)+1;period=f"{f['year']}-{month:02}"
  # Redundant header anchors recover fichas with a missing DADOS cell.
  starts=set()
  for a,e in c.items():
   col=re.sub(r'\d','',a);row=int(re.search(r'\d+',a)[0]);v=norm(e['v'])
   if col=='B' and v=='dados' or col=='C' and v=='cliente':starts.add(row)
  starts=sorted(starts);sheetrecs=[];blank=0;unassigned=[]
  for idx,start in enumerate(starts):
   end=min(starts[idx+1] if idx+1<len(starts) else start+19,start+38)
   raw={a:v for a,v in c.items() if re.fullmatch(r'[A-O]\d+',a) and start<=int(re.search(r'\d+',a)[0])<end}
   get=lambda col,r:value(c,f'{col}{r}')
   meta={k:get('C',start+offset) for offset,k in [(1,'date'),(2,'name'),(3,'plate'),(4,'model'),(5,'km')]}
   if not text(meta['date']) and re.fullmatch(r'DATA\s+\d{1,2}/\d{1,2}/\d{2,4}',text(get('B',start+1)),re.I):meta['date']=re.sub(r'^DATA\s+','',text(get('B',start+1)),flags=re.I)
   # Exclude links used to navigate to other workbooks; retain them in raw cells.
   items=[]
   for r in range(start+1,end):
    ds,pt,qt,vl=[get(col,r) for col in ['E','G','I','K']]
    if any('http' in text(x).lower() for x in [ds,pt]):continue
    if not any(text(x) for x in [ds,pt,qt,vl]):continue
    if not text(pt) and not text(qt) and not text(vl) and (len(text(ds))<2 or text(ds).isdigit()):continue
    items.append({'row':r,'description':text(ds) or None,'part':text(pt) or None,'quantity':text(qt) or None,'value':text(vl) or None,'valueCents':money(vl)})
   meaningful=any(text(x) for x in meta.values()) or bool(items) or any(money(get('N',r)) not in (None,0) for r in range(start+1,end))
   if not meaningful:
    isolated=[{'cell':f'{col}{r}','value':text(get(col,r)),'valueCents':money(get(col,r))} for r in range(start+1,end) for col in ['M','O'] if money(get(col,r)) not in (None,0)]
    if isolated:unassigned.append({'row':start,'values':isolated,'note':'Valor isolado em ficha vazia; preservado para conferência, sem criar atendimento.'})
   if not meaningful:blank+=1;continue
   issues=[];date,dateend,precision,suggestion=parse_date(meta['date'],f['year'],month)
   if precision=='month':issues.append('Data ausente ou divergente da aba; o mês é a referência do arquivo e o dia precisa de conferência.')
   customer=text(meta['name']);model=text(meta['model']);plate=text(meta['plate']);canonical=re.sub('[^A-Z0-9]','',plate.upper())
   if not re.fullmatch(r'[A-Z]{3}[0-9][A-Z0-9][0-9]{2}',canonical):canonical=None;issues.append('Placa ausente ou fora do padrão; não foi corrigida automaticamente.')
   if not customer:issues.append('Nome do cliente não informado na ficha.')
   if not model:issues.append('Modelo da moto não informado na ficha.')
   km=meta['km'];odometer=None
   if isinstance(km,(int,float)) and km==int(km) and 0<=km<=9999999:odometer=int(km)
   elif re.fullmatch(r'\d{1,7}',text(km)):odometer=int(text(km))
   elif text(km):issues.append('Quilometragem não reconhecida; consulte o valor original.')
   partstotal=labortotal=grand=None;grandcell=None;partscell=None;laborcell=None
   for r in range(start+1,end):
    label=norm(get('M',r))
    if label in ['valor total','valor pecas']:partstotal=money(get('O',r));partscell=f'O{r}'
    if label.startswith('m.o'):
     col='N' if get('N',r) is not None else 'O';labortotal=money(get(col,r));laborcell=f'{col}{r}'
    if label=='total geral':grand=money(get('M',r+1));grandcell=f'M{r+1}'
   subtotal=sum(i['valueCents'] or 0 for i in items);calculated=(partstotal or 0)+(labortotal or 0)
   basis='stated' if grand is not None else 'unknown'
   # Never replace a missing grand total with a guessed zero or sum.
   if grand is None:issues.append('Total geral não informado ou fórmula sem valor válido; confira os valores originais.')
   if grand is not None and (partstotal is not None or labortotal is not None) and grand!=calculated:issues.append('Total geral diverge da soma de peças e mão de obra; mantido o total da planilha.')
   if partstotal is not None and partstotal!=subtotal:issues.append('Total de peças diverge da soma da coluna VALOR; mantidos os dois valores originais.')
   if not items:issues.append('Descrição de serviços e peças não informada.')
   note=[];phone=''
   label6=norm(get('B',start+6))
   if label6 in ['tel','tel;','fone']:phone=text(get('C',start+6))
   elif text(get('C',start+6)):note.append(text(get('C',start+6)))
   for r in range(start+7,end):
    for col in ['B','C','D']:
     v=text(get(col,r))
     if v:note.append(v)
   record={'sheet':name,'row':start,'endRow':end-1,'sourceNumber':text(get('A',start)),'date':date,'dateEnd':dateend,'datePrecision':precision,'dateOriginal':text(meta['date']),'suggestedDate':suggestion,'name':customer or None,'phone':phone,'plate':plate or None,'canonicalPlate':canonical,'model':model or None,'odometer':odometer,'service':' · '.join(' / '.join(v for v in [i['description'],i['part']] if v) for i in items if i['description'] or i['part']) or 'Descrição não informada na planilha','items':items,'partsCents':partstotal,'laborCents':labortotal,'totalCents':grand,'totalBasis':basis,'totalCell':grandcell,'partsCell':partscell,'laborCell':laborcell,'notes':'\n'.join(note),'issues':issues,'raw':raw}
   sheetrecs.append(record)
  expenses=[]
  for a,e in c.items():
   if re.fullmatch('Q[0-9]+',a) and text(e['v']) and norm(e['v'])!='despesa':
    r=int(a[1:]);expenses.append({'row':r,'description':text(e['v']),'value':text(value(c,'S'+str(r))),'valueCents':money(value(c,'S'+str(r)))})
  summaries=[]
  for a,e in c.items():
   if re.fullmatch('U[0-9]+',a) and text(e['v']) and norm(e['v'])!='info mensal':
    r=int(a[1:]);summaries.append({'cell':'W'+str(r),'label':text(e['v']),'value':text(value(c,'W'+str(r))),'valueCents':money(value(c,'W'+str(r)))})
  sheet={'name':name,'kind':'month','period':period,'records':len(sheetrecs),'blankForms':blank,'unassignedTotals':unassigned,'totalCents':sum(r['totalCents'] or 0 for r in sheetrecs),'unknownTotals':sum(r['totalCents'] is None for r in sheetrecs),'datesToReview':sum(r['datePrecision']=='month' for r in sheetrecs),'expenses':expenses,'summaries':summaries,'nonemptyCells':len(c)}
  comparisons=[]
  for summary in summaries:
   label=norm(summary['label'])
   extracted=None
   if label=='entrada pecas':extracted=sum(r['partsCents'] or 0 for r in sheetrecs)
   if label=='total m.o':extracted=sum(r['laborCents'] or 0 for r in sheetrecs)
   if label=='despesas total':extracted=sum(e['valueCents'] or 0 for e in expenses)
   if extracted is not None:comparisons.append({'label':summary['label'],'sourceCell':summary['cell'],'sourceCents':summary['valueCents'],'extractedCents':extracted,'agrees':extracted==summary['valueCents']})
  sheet['reconciliation']=comparisons
  records+=sheetrecs;sheets.append(sheet)
 records.sort(key=lambda r:(MONTHS.index(r['sheet']),r['row']))
 # Identical content at different source locations is retained and flagged, never silently removed.
 seen={}
 for r in records:
  signature=hashlib.sha256(json.dumps([r[k] for k in ['dateOriginal','name','plate','model','odometer','service','totalCents']],ensure_ascii=False).encode()).hexdigest()
  if signature in seen:
   previous=seen[signature];message=f"Possível repetição de {previous['sheet']}, linha {previous['row']}; ficha preservada."
   r['issues'].append(message)
  else:seen[signature]=r
 return {'version':'oficina-excel-v1','filename':f['filename'],'fileHash':hashlib.sha256((folder/f['filename']).read_bytes()).hexdigest(),'year':f['year'],'records':records,'sheets':sheets}

def main():
 a=argparse.ArgumentParser();a.add_argument('input',type=Path);a.add_argument('output',type=Path);a.add_argument('--cache',type=Path);args=a.parse_args();args.output.mkdir(parents=True,exist_ok=True)
 originals=json.loads(args.cache.read_text()) if args.cache else read_workbooks(args.input)
 reports=[]
 for f in sorted(originals,key=lambda f:f['year']):
  manifest=prepare(f,args.input)
  (args.output/f"{f['year']}.json").write_text(json.dumps(manifest,ensure_ascii=False,separators=(',',':'),default=str))
  report={'year':f['year'],'filename':f['filename'],'records':len(manifest['records']),'datesToReview':sum(r['datePrecision']=='month' for r in manifest['records']),'totalCents':sum(r['totalCents'] or 0 for r in manifest['records']),'unknownTotals':sum(r['totalCents'] is None for r in manifest['records']),'issues':dict(collections.Counter(i for r in manifest['records'] for i in r['issues']))}
  reports.append(report);print(json.dumps({k:v for k,v in report.items() if k!='issues'},ensure_ascii=False))
 (args.output/'reconciliation.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
